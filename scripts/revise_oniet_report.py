from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(r'C:\Users\marce\Desktop\codi\Informe_Evaluador_Automatico_EVA-3117_y_Codi.docx')
OUTPUT = Path(
    r'C:\Users\marce\Desktop\codi\Informe_Evaluador_Automatico_EVA-3117_y_Codi_final_completo.docx'
)
FIELD_IMAGES = Path(r'C:\Users\marce\Documents\Codi\architecture\images\carpeta de campo')
ROOTS_EXERCISE = Path(r'C:\Users\marce\Desktop\evaluador\Obteniendo Raices')
TEMP_DIR = Path(r'C:\Users\marce\AppData\Local\Temp\opencode')

FIRST_SURVEY_IMAGES = [
    'encuesta (8).jpeg',
    'encuesta (1).jpeg',
    'encuesta (7).jpeg',
    'encuesta (13).jpeg',
    'encuesta (6).jpeg',
    'encuesta (12).jpeg',
    'encuesta (5).jpeg',
    'encuesta (11).jpeg',
    'encuesta (4).jpeg',
    'encuesta (10).jpeg',
    'encuesta (3).jpeg',
    'encuesta (9).jpeg',
    'encuesta (2).jpeg',
    'encuesta (14).jpeg',
    'encuesta.jpeg',
]


REPLACEMENTS = {
    'EVA 3117 comenzó el 7 de mayo de 2026 como respuesta a la necesidad de disponer de un juez y entorno '
    'de concursos para la escuela. La decisión técnica fue adoptar CMS como base de código abierto y realizar '
    'una implementación institucional. Codi comenzó el 7 de julio de 2026, dos meses después, para cubrir '
    'una necesidad complementaria: no solo evaluar una entrega, sino organizar la enseñanza previa, la práctica '
    'y el seguimiento. La primera prueba con usuarios se realizó el 9 de junio de 2026 en el laboratorio de '
    'Informática; el primer simulacro institucional se registró el 12 de junio de 2026.':
        'EVA 3117 comenzó el 7 de mayo de 2026 como respuesta a la necesidad de disponer de un juez y entorno '
        'de concursos para la escuela. La decisión técnica fue adoptar CMS como base de código abierto y realizar '
        'una implementación institucional. La primera prueba con usuarios de EVA 3117 se realizó el 9 de junio '
        'de 2026 en el laboratorio de Informática y el primer simulacro institucional se registró el 12 de junio. '
        'Codi comenzó el 7 de julio de 2026 para cubrir una necesidad complementaria: no solo evaluar una entrega, '
        'sino organizar la enseñanza previa, la práctica y el seguimiento. Por lo tanto, la prueba de junio valida '
        'EVA 3117, mientras que la experiencia integrada con Codi corresponde a las instancias posteriores.',
    'El 9 de junio de 2026 se realizó una primera prueba con usuarios en el laboratorio de Informática. Los participantes '
    'realizaron varios intentos sobre las actividades propuestas y recibieron respuestas en menos de cuatro segundos tanto '
    'desde Codi como desde EVA 3117. Posteriormente, el 23 de agosto de 2026 se realizó un examen previo con 21 estudiantes. '
    'Durante esa instancia todos los envíos fueron procesados correctamente y no se observaron pérdidas de entregas ni fallas '
    'funcionales. La única limitación detectada fue de rendimiento: al aumentar la cantidad de usuarios que enviaban soluciones '
    'al mismo tiempo también aumentó el tiempo de espera para procesar la cola y devolver cada resultado. Este comportamiento '
    'orienta la optimización futura hacia pruebas de carga escalonadas, medición de latencia y ampliación de recursos cuando sea '
    'necesario.':
        'El 9 de junio de 2026 se realizó una primera prueba con usuarios de EVA 3117 en el laboratorio de '
        'Informática. Los participantes realizaron varios intentos sobre las actividades propuestas y recibieron '
        'respuestas en menos de cuatro segundos. Codi aún no había comenzado, por lo que esta instancia no se '
        'presenta como una prueba de la integración. Posteriormente, el 23 de agosto de 2026 se realizó un examen '
        'previo con 21 estudiantes. Durante esa instancia todos los envíos fueron procesados correctamente y no se '
        'observaron pérdidas de entregas ni fallas funcionales. La limitación detectada fue de rendimiento: al '
        'aumentar los envíos simultáneos también aumentó el tiempo de espera para procesar la cola y devolver cada '
        'resultado. Este comportamiento orienta la optimización futura hacia pruebas de carga escalonadas, medición '
        'de latencia y ampliación de recursos cuando sea necesario.',
    'La innovación central es articular aprendizaje y evaluación dentro de un solo sistema. Codi presenta el '
    'recorrido de aprendizaje, contenidos, progreso, laboratorio, recompensas y ranking; EVA 3117 administra '
    'concursos, tareas, casos de prueba, graders y veredictos. EVA se construye sobre CMS; Codi fue desarrollado '
    'por el equipo con Next.js y NestJS para integrarse con EVA. Comparten infraestructura y datos: usuarios, '
    'actividades, envíos y resultados se mantienen sincronizados. La diferencia no se limita a la apariencia: '
    'transforma la forma de organizar y utilizar la información de un juez de competencias para adecuarla al '
    'aprendizaje de estudiantes de una escuela secundaria.':
        'La innovación central es articular aprendizaje y evaluación dentro de un mismo ecosistema. Codi presenta '
        'recorridos, contenidos, progreso, laboratorio, recompensas y ranking; EVA 3117 aporta la configuración de '
        'concursos, tareas, datasets, casos de prueba y graders de CMS. Ambos componentes utilizan una única base '
        'PostgreSQL: Codi autentica cuentas de CMS y consulta sus tareas y datos de evaluación, mientras conserva '
        'progreso, recompensas y solicitudes de evaluación en tablas propias con prefijo codi_. Cada problema de '
        'Codi referencia la tarea correspondiente de CMS mediante su identificador. La API coloca el envío en una '
        'cola persistida, el worker obtiene la configuración de CMS, ejecuta Python o C++ en Cloudflare Sandbox y '
        'registra el resultado para mostrarlo en Codi. La integración se basa así en referencias y datos compartidos, '
        'no en dos copias sincronizadas ni en una modificación que atribuya al equipo la autoría de CMS.',
}


HEADING_RENAMES = {
    '7.5 Guía operativa: carga de problemas y uso estudiantil': 'Anexo A. Guía operativa: carga de problemas y uso estudiantil',
    '7.5.1 Paso a paso para cargar un problema': 'A.1 Paso a paso para cargar un problema',
    '7.5.2 Cómo se construyen las entradas y salidas': 'A.2 Cómo se construyen las entradas y salidas',
    '7.5.3 Guía del estudiante': 'A.3 Guía del estudiante',
    '9.4 Registro y resultados de la encuesta': 'Anexo B. Registro y resultados de las encuestas',
    '9.4.1 Registro de participantes e instrumentos': 'B.1 Registro de participantes e instrumentos',
    '9.4.2 Caracterización general de los estudiantes': 'B.2 Caracterización general de los estudiantes',
    '9.4.3 Experiencia de uso de Codi': 'B.3 Experiencia de uso de Codi',
    '9.4.4 Síntesis de las preguntas abiertas': 'B.4 Síntesis de las preguntas abiertas',
}


def replace_paragraph_text(document: Document) -> None:
    paragraphs_by_text = {paragraph.text: paragraph for paragraph in document.paragraphs}
    for old_text, new_text in REPLACEMENTS.items():
        paragraph = paragraphs_by_text.get(old_text)
        if paragraph is None:
            raise ValueError(f'Required paragraph was not found: {old_text[:80]}')
        paragraph.text = new_text


def rename_annex_headings(document: Document) -> None:
    for paragraph in document.paragraphs:
        renamed = HEADING_RENAMES.get(paragraph.text)
        if renamed is None:
            continue
        paragraph.text = renamed
        paragraph.style = 'Heading 1' if renamed.startswith('Anexo ') else 'Heading 2'


def relabel_annex_figures(document: Document) -> None:
    for paragraph in document.paragraphs:
        for number in range(25, 33):
            prefix = f'Figura {number}.'
            if paragraph.text.startswith(prefix):
                paragraph.text = paragraph.text.replace(prefix, f'Figura A.{number - 24}.', 1)
        for number in range(34, 48):
            prefix = f'Figura {number}.'
            if paragraph.text.startswith(prefix):
                paragraph.text = paragraph.text.replace(prefix, f'Figura B.{number - 33}.', 1)
        if paragraph.text.startswith('Figura 33.'):
            paragraph.text = paragraph.text.replace('Figura 33.', 'Figura 25.', 1)


def move_range_to_end(document: Document, start_text: str, end_text: str) -> None:
    body = document.element.body
    start = next(paragraph._p for paragraph in document.paragraphs if paragraph.text == start_text)
    end = next(paragraph._p for paragraph in document.paragraphs if paragraph.text == end_text)
    children = list(body)
    moving = children[children.index(start) : children.index(end)]
    section_properties = body.find(qn('w:sectPr'))

    for child in moving:
        body.remove(child)
        section_properties.addprevious(child)


def request_field_update(document: Document) -> None:
    settings = document.settings.element
    update_fields = settings.find(qn('w:updateFields'))
    if update_fields is None:
        update_fields = OxmlElement('w:updateFields')
        settings.append(update_fields)
    update_fields.set(qn('w:val'), 'true')


def add_centered_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def add_code_block(document: Document, source: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.right_indent = Inches(0.3)
    paragraph.paragraph_format.space_after = Pt(8)
    for line_number, line in enumerate(source.splitlines(), start=1):
        run = paragraph.add_run(f'{line_number:>2}  {line}\n')
        run.font.name = 'Consolas'
        run.font.size = Pt(8)


def read_test_cases() -> list[tuple[str, str, str, str]]:
    input_lines = (ROOTS_EXERCISE / 'in.txt').read_text(encoding='utf-8').splitlines()
    output_lines = (ROOTS_EXERCISE / 'out.txt').read_text(encoding='utf-8').splitlines()
    cases = []
    for index in range(0, len(input_lines), 2):
        case_number = input_lines[index].removesuffix('.in')
        case_input = input_lines[index + 1]
        case_output = output_lines[index + 1]
        subtask = '1. Raíces reales' if case_output != '-1' else '2. Sin raíces reales'
        cases.append((case_number, case_input, case_output, subtask))
    return cases


def add_test_case_table(document: Document, table_style) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = table_style
    headers = ['Caso', 'Entrada: A B C', 'Salida esperada', 'Subtarea']
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for case in read_test_cases():
        cells = table.add_row().cells
        for cell, text in zip(cells, case):
            cell.text = text
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def add_first_survey_annex(document: Document, table_style) -> None:
    document.add_page_break()
    document.add_heading('Anexo C. Encuesta N.° 1 sobre EVA 3117', level=1)
    document.add_paragraph(
        'Primera encuesta sobre tiempos de devolución, corrección automática y expectativas respecto de EVA 3117. '
        'Se incorporan los 15 gráficos originales en el orden numérico de las preguntas. La cantidad de respuestas '
        'visible en cada captura se conserva sin modificaciones.'
    )

    for first_question in range(0, len(FIRST_SURVEY_IMAGES), 4):
        table = document.add_table(rows=2, cols=2)
        table.style = table_style
        page_images = FIRST_SURVEY_IMAGES[first_question : first_question + 4]
        for offset, image_name in enumerate(page_images):
            question_number = first_question + offset + 1
            cell = table.cell(offset // 2, offset % 2)
            label = cell.paragraphs[0]
            label.alignment = WD_ALIGN_PARAGRAPH.CENTER
            label.add_run(f'Pregunta {question_number}').bold = True
            picture = cell.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.add_run().add_picture(str(FIELD_IMAGES / image_name), width=Inches(2.85))
        if first_question + 4 < len(FIRST_SURVEY_IMAGES):
            document.add_page_break()


def render_roots_statement() -> Path:
    statement_pdf = ROOTS_EXERCISE / 'cuadratica.pdf'
    statement_image = TEMP_DIR / 'obteniendo-raices.png'
    with pymupdf.open(statement_pdf) as pdf:
        page = pdf[0]
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
        pixmap.save(statement_image)
    return statement_image


def add_roots_exercise_annex(document: Document, table_style) -> None:
    document.add_page_break()
    document.add_heading('Anexo D. Primer ejercicio: Obteniendo raíces', level=1)
    document.add_paragraph(
        'Enunciado original del primer ejercicio solicitado. El material fuente se encuentra en '
        'Desktop\\evaluador\\Obteniendo Raices e incluye cuadratica.pdf, cuadratica.py, solution.py, in.txt y '
        'out.txt. Los archivos de entrada y salida reúnen 40 casos de prueba.'
    )
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(render_roots_statement()), width=Inches(6.25))
    add_centered_caption(document, 'Figura D.1. Enunciado original del ejercicio Obteniendo raíces.')

    document.add_heading('D.1 Plantilla entregada al estudiante', level=2)
    document.add_paragraph(
        'Contenido exacto de solution.py. Se conserva tal como fue suministrado en la carpeta del ejercicio.'
    )
    add_code_block(document, (ROOTS_EXERCISE / 'solution.py').read_text(encoding='utf-8'))

    document.add_heading('D.2 Solución de referencia del evaluador', level=2)
    document.add_paragraph(
        'Contenido exacto de cuadratica.py, utilizado como referencia para calcular las raíces y producir la salida esperada.'
    )
    add_code_block(document, (ROOTS_EXERCISE / 'cuadratica.py').read_text(encoding='utf-8'))

    document.add_heading('D.3 Subtareas y puntuación', level=2)
    document.add_paragraph(
        'Subtarea 1 (50 puntos): retornar correctamente x1 y x2, ordenadas de menor a mayor, cuando el '
        'discriminante es mayor o igual que cero y existen raíces reales enteras.'
    )
    document.add_paragraph(
        'Subtarea 2 (50 puntos): retornar -1 cuando el discriminante es negativo y el polinomio no tiene raíces reales.'
    )

    document.add_heading('D.4 Datos de entrada y salida', level=2)
    document.add_paragraph(
        'La tabla transcribe los 40 pares originales de in.txt y out.txt. Los casos 1 a 30 verifican raíces reales '
        'y los casos 31 a 40 verifican polinomios sin raíces reales.'
    )
    add_test_case_table(document, table_style)


def revise_report() -> None:
    document = Document(SOURCE)
    table_style = document.tables[0].style
    replace_paragraph_text(document)
    rename_annex_headings(document)
    relabel_annex_figures(document)
    move_range_to_end(document, 'Anexo A. Guía operativa: carga de problemas y uso estudiantil', '8. Impacto esperado')
    move_range_to_end(document, 'Anexo B. Registro y resultados de las encuestas', '9.5 Mejoras previstas a partir del relevamiento')
    add_first_survey_annex(document, table_style)
    add_roots_exercise_annex(document, table_style)
    move_range_to_end(document, '10. Bibliografía consultada', 'Anexo A. Guía operativa: carga de problemas y uso estudiantil')
    bibliography = next(paragraph for paragraph in document.paragraphs if paragraph.text == '10. Bibliografía consultada')
    bibliography.paragraph_format.page_break_before = True
    request_field_update(document)
    document.save(OUTPUT)


if __name__ == '__main__':
    revise_report()
    print(OUTPUT)
